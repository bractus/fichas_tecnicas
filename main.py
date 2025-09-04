import os
import json
import logging
import warnings
import tracemalloc
from datetime import datetime
from typing import List, Optional
#from dotenv import load_dotenv

# Enable tracemalloc for better resource tracking
tracemalloc.start()

# Suppress specific warnings
warnings.filterwarnings("ignore", message="Field name \"json\" in \"ChangeTrackingData\" shadows an attribute in parent \"BaseModel\"")
warnings.filterwarnings("ignore", category=DeprecationWarning, module="alembic")
warnings.filterwarnings("ignore", category=ResourceWarning, module="urllib3")
warnings.filterwarnings("ignore", message="unclosed.*ssl.SSLSocket.*")
warnings.filterwarnings("ignore", category=ResourceWarning, message="unclosed.*ssl.SSLSocket.*")
from crewai import Agent, Task, Crew, Process
from crewai import LLM
import pandas as pd
from docx import Document
import PyPDF2
from pathlib import Path
from urllib.parse import urlparse
from pydantic import BaseModel, Field, ConfigDict, validator
from enum import Enum
import yaml
import shutil
from crewai_tools import EXASearchTool
from crewai_tools import SerperDevTool, FileWriterTool, RagTool
from tools.webscraping import WebScrapingTool
from tools.file_reader import MultiFormatFileReader
from tools.generate_excel import ExcelGeneratorTool
from datetime import datetime

from openinference.instrumentation.crewai import CrewAIInstrumentor
from openinference.instrumentation.litellm import LiteLLMInstrumentor


import chromadb

# Use the new ChromaDB API
client = chromadb.PersistentClient(path=".chromadb")

# --- PYDANTIC MODELS ---
class UnidadeMedida(str, Enum):
    KG = "kg"
    L = "L"

class SourceContent(BaseModel):
    fonte: str = Field(..., description="Identificação da fonte (arquivo ou URL)")
    conteudo: str = Field(..., description="Conteúdo extraído da fonte")
    tipo: str = Field(..., description="Tipo da fonte (arquivo, URL)")
    status: str = Field(..., description="Status da extração (sucesso, erro)")

class Ingrediente(BaseModel):
    model_config = ConfigDict(extra="allow")
    
    nome: str = Field(..., description="Nome do ingrediente")
    unidade: UnidadeMedida = Field(..., description="Unidade de medida obrigatoriamente em Kg ou L")
    quantidade: float = Field(..., description="Quantidade necessária")
    fator_correcao: float = Field(..., description="Fator de correção do ingrediente (ex: 1.2 para 20% de perda)")
    custo_unitario: float = Field(..., description="Custo unitário do ingrediente")

class Insumo(BaseModel):
    model_config = ConfigDict(extra="allow")
    
    ingrediente: str = Field(..., description="Nome do ingrediente")
    unidade: UnidadeMedida = Field(..., description="Unidade de medida obrigatoriamente em Kg ou L")
    preco: float = Field(..., description="Preço unitário")
    fator_correcao: float = Field(..., description="Fator de correção do ingrediente")
    fornecedor: str = Field(..., description="Fornecedor sugerido")
    data_cotacao: str = Field(..., description="Data da cotação")

class FichaTecnica(BaseModel):
    model_config = ConfigDict(extra="allow")
    
    nome_preparacao: str = Field(..., description="Nome da receita")
    rendimento_porcoes: int = Field(..., description="Número de porções")
    preco_venda: float = Field(..., description="Preço de venda por porção (obrigatório)")
    cmv_percentual: Optional[float] = Field(None, description="CMV (Custo da Mercadoria Vendida) em percentual")
    ingredientes: List[Ingrediente] = Field(..., description="Lista de ingredientes")
    modo_preparo: List[str] = Field(default=[], description="Lista de passos do modo de preparo")

class FichaTecnicaList(BaseModel):
    model_config = ConfigDict(extra="allow")
    
    fichas_tecnicas: List[FichaTecnica] = Field(..., description="Lista de fichas técnicas extraídas")

class InsumoList(BaseModel):
    model_config = ConfigDict(extra="allow")
    
    insumos: List[Insumo] = Field(..., description="Lista de insumos consolidados")

class ExcelGenerationResult(BaseModel):
    arquivo_gerado: str = Field(..., description="Caminho completo do arquivo Excel gerado")
    numero_fichas: int = Field(..., description="Número de fichas técnicas processadas")
    numero_insumos: int = Field(..., description="Número de insumos processados")
    status: str = Field(..., description="Status da geração (sucesso, erro)")
    timestamp: str = Field(..., description="Timestamp da geração")

class RecipeData(BaseModel):
    fichas_tecnicas: List[FichaTecnica] = Field(..., description="Lista de todas as fichas técnicas")
    base_de_insumos: List[Insumo] = Field(..., description="Base unificada de insumos")


from langfuse import Langfuse

# Configure logging with proper file handle management
log_file_handler = logging.FileHandler('fichas_tecnicas.log')
stream_handler = logging.StreamHandler()

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[log_file_handler, stream_handler]
)
logger = logging.getLogger(__name__)

def init_langfuse():
    """Initialize Langfuse client with credentials from environment."""
    try:
        langfuse = Langfuse(
            secret_key=os.getenv("LANGFUSE_SECRET_KEY"),
            public_key=os.getenv("LANGFUSE_PUBLIC_KEY"),
            host=os.getenv("LANGFUSE_HOST", "https://us.cloud.langfuse.com")
        )
        
        if langfuse.auth_check():
            logger.info("Langfuse client authenticated successfully")
            return langfuse
        else:
            logger.warning("Langfuse authentication failed")
            return None
    except Exception as e:
        logger.error(f"Failed to initialize Langfuse: {e}")
        return None

 
CrewAIInstrumentor().instrument(skip_dep_check=True)
LiteLLMInstrumentor().instrument()


def validate_recipe_data(json_data: str) -> RecipeData:
    """Valida e converte JSON para o modelo Pydantic RecipeData."""
    try:
        data_dict = json.loads(json_data) if isinstance(json_data, str) else json_data
        
        # Log validation details
        logger.info("=== VALIDAÇÃO DE DADOS ===")
        logger.info(f"Tipo dos dados recebidos: {type(data_dict)}")
        logger.info(f"Chaves no JSON: {list(data_dict.keys()) if isinstance(data_dict, dict) else 'Não é dicionário'}")
        
        if isinstance(data_dict, dict):
            fichas_count = len(data_dict.get('fichas_tecnicas', []))
            insumos_count = len(data_dict.get('base_de_insumos', []))
            logger.info(f"Fichas técnicas encontradas: {fichas_count}")
            logger.info(f"Insumos encontrados: {insumos_count}")
            
            # Log individual recipe names for debugging
            if 'fichas_tecnicas' in data_dict:
                for i, ficha in enumerate(data_dict['fichas_tecnicas'], 1):
                    nome = ficha.get('nome_preparacao', 'NOME AUSENTE') if isinstance(ficha, dict) else 'FORMATO INVÁLIDO'
                    logger.info(f"Ficha {i} na validação: {nome}")
        
        result = RecipeData(**data_dict)
        logger.info(f"✅ Validação bem-sucedida: {len(result.fichas_tecnicas)} fichas, {len(result.base_de_insumos)} insumos")
        return result
        
    except Exception as e:
        logger.error(f"❌ Erro na validação dos dados: {e}")
        logger.error(f"Dados que causaram erro: {str(json_data)[:500]}...")
        raise


def ensure_output_directory() -> str:
    """Garante que o diretório output existe e retorna o caminho absoluto."""
    current_working_dir = os.getcwd()
    logger.info(f"Current working directory: {current_working_dir}")
    
    # Use hardcoded project path to ensure correct location
    project_root = "/Users/cairorocha/Documents/fichas_tecnicas1"
    
    # Verify the project directory exists
    if os.path.exists(project_root):
        output_dir = os.path.join(project_root, "output")
        logger.info(f"Using hardcoded project root: {project_root}")
    else:
        # Fallback to calculated path if hardcoded doesn't exist
        project_root = os.path.dirname(os.path.abspath(__file__))
        output_dir = os.path.join(project_root, "output")
        logger.info(f"Fallback to calculated project root: {project_root}")
    
    logger.info(f"Output directory ensured: {output_dir}")
    os.makedirs(output_dir, exist_ok=True)
    return output_dir


def load_config(config_file: str) -> dict:
    """Carrega configurações de arquivos YAML."""
    config_path = os.path.join("config", config_file)
    try:
        with open(config_path, 'r', encoding='utf-8') as file:
            config = yaml.safe_load(file)
            logger.info(f"Configuration loaded: {config_file}")
            return config
    except FileNotFoundError:
        logger.error(f"Configuration file not found: {config_path}")
        raise
    except yaml.YAMLError as e:
        logger.error(f"Error parsing YAML file {config_file}: {e}")
        raise


def create_agent_from_config(agent_name: str, agent_config: dict, tools_dict: dict, llm) -> Agent:
    """Cria um agente a partir da configuração YAML."""
    tools = [tools_dict[tool_name] for tool_name in agent_config['tools']]
    
    return Agent(
        role=agent_config['role'],
        goal=agent_config['goal'],
        backstory=agent_config['backstory'],
        tools=tools,
        verbose=agent_config.get('verbose', True),
        llm=llm
    )


def create_task_from_config(task_name: str, task_config: dict, agents_dict: dict, context_tasks: dict = None, **kwargs) -> Task:
    """Cria uma task a partir da configuração YAML."""
    agent = agents_dict[task_config['agent']]
    
    # Substituir placeholders na descrição
    description = task_config['description']
    for key, value in kwargs.items():
        description = description.replace(f"{{{key}}}", str(value))
    
    task_kwargs = {
        'agent': agent,
        'description': description,
        'expected_output': task_config['expected_output']
    }
    
    # Adicionar contexto se especificado
    if 'context' in task_config and context_tasks:
        task_kwargs['context'] = [context_tasks[ctx] for ctx in task_config['context']]
    
    # Adicionar output_file se especificado
    if 'output_file' in task_config:
        output_file = task_config['output_file']
        for key, value in kwargs.items():
            output_file = output_file.replace(f"{{{key}}}", str(value))
        task_kwargs['output_file'] = output_file
    
    return Task(**task_kwargs)


def validate_sources(sources: List[str]) -> List[str]:
    """Valida e filtra fontes válidas."""
    sources_validas = []
    for source in sources:
        if source.startswith(('http://', 'https://')):
            try:
                parsed = urlparse(source)
                if parsed.netloc and parsed.scheme:
                    sources_validas.append(source)
                    logger.info(f"Valid URL added: {source}")
                else:
                    logger.warning(f"Invalid URL format: {source}")
            except Exception as e:
                logger.warning(f"URL validation failed for {source}: {e}")
        elif os.path.exists(source):
            sources_validas.append(source)
            logger.info(f"Valid file found: {source}")
        else:
            logger.warning(f"Source not found: {source}")
    
    return sources_validas

def safe_file_check(file_path: str) -> bool:
    """Safely check if file is valid Excel format."""
    try:
        if not os.path.exists(file_path):
            return False
        
        # Use file extension as primary check
        if file_path.lower().endswith('.xlsx'):
            # Check file size (should be reasonable minimum)
            size = os.path.getsize(file_path)
            if size < 1024:  # Less than 1KB is probably invalid
                return False
                
            # Try to open with openpyxl to verify it's a valid Excel file
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path)
                # Check if it has worksheets
                if len(wb.worksheets) == 0:
                    return False
                # Try to access first sheet
                first_sheet = wb.worksheets[0]
                # Try to read a cell
                _ = first_sheet.cell(1, 1).value
                wb.close()
                return True
            except Exception:
                return False
        return False
    except Exception as e:
        logger.error(f"Error checking file {file_path}: {e}")
        return False

def fichas_tecnicas(custom_sources=None):
    """Gera fichas técnicas de receitas culinárias a partir de múltiplos arquivos/URLs.
    
    Args:
        custom_sources (list, optional): Lista de fontes customizadas para processar.
                                       Se None, usa sources.yaml padrão.
    """
    try:
        # Carregar variáveis de ambiente
        #load_dotenv()

        logger.info("Starting fichas técnicas generation process")

        # --- Configuração do LLM ---
        api_key = os.getenv("OPENAI_API_KEY")
        if not api_key:
            raise ValueError("OPENAI_API_KEY not found in environment variables")

        llm = LLM(
            model='openai/gpt-4.1-nano',
            api_key=api_key,
            temperature=0.0
        )
        llm2 = LLM(
            model='gemini/gemini-2.5-flash',
            api_key=os.getenv("GEMINI_API_KEY"),
            temperature=0.0
        )

        llm3 = LLM(
            model='openai/gpt-4o-mini',
            api_key=api_key,
            temperature=0.0
        )

        logger.info("LLM models configured successfully")

        # --- FERRAMENTAS ---
        try:
            web_search = EXASearchTool()
            multi_reader_tool = MultiFormatFileReader()
            excel_generator_tool = ExcelGeneratorTool()
            web_scraping_tool = WebScrapingTool()
            file_writer_tool = FileWriterTool()

            # RAG Tool for correction factors
            fatores_rag = RagTool(
                name="Fatores de Correção RAG",
                description="Base de conhecimento com fatores de correção para ingredientes culinários. Use para consultar fatores de correção antes de buscar na web."
            )
            # Add the markdown files to the RAG tool
            fatores_rag.add("./files/fc/fatores1.md")
            fatores_rag.add("./files/fc/fatores2.md")
            
            logger.info("Tools initialized successfully")
        except Exception as e:
            logger.error(f"Error initializing tools: {e}")
            raise

        tools_dict = {
            'web_search': web_search,
            'multi_reader_tool': multi_reader_tool,
            'excel_generator_tool': excel_generator_tool,
            'file_writer_tool': file_writer_tool,
            'web_scraping_tool': web_scraping_tool,
            'fatores_rag': fatores_rag
        }

        # --- CARREGAR CONFIGURAÇÕES ---
        agents_config = load_config('agents.yaml')
        tasks_config = load_config('tasks.yaml')
        sources_config = load_config('sources.yaml')

        # --- CRIAR AGENTES A PARTIR DOS CONFIGS ---
        try:
            file_reader_agent = create_agent_from_config('file_reader_agent', agents_config['file_reader_agent'], tools_dict, llm2)
            ficha_tecnica_agent = create_agent_from_config('ficha_tecnica_agent', agents_config['ficha_tecnica_agent'], tools_dict, llm3)
            base_insumos_agent = create_agent_from_config('base_insumos_agent', agents_config['base_insumos_agent'], tools_dict, llm3)
            data_consolidator_agent = create_agent_from_config('data_consolidator_agent', agents_config['data_consolidator_agent'], tools_dict, llm)
            excel_writer_agent = create_agent_from_config('excel_writer_agent', agents_config['excel_writer_agent'], tools_dict, llm2)
            logger.info("Agents created successfully")
        except Exception as e:
            logger.error(f"Error creating agents: {e}")
            raise

        agents_dict = {
            'file_reader_agent': file_reader_agent,
            'ficha_tecnica_agent': ficha_tecnica_agent,
            'base_insumos_agent': base_insumos_agent,
            'data_consolidator_agent': data_consolidator_agent,
            'excel_writer_agent': excel_writer_agent
        }

        # --- DEFINIR VARIÁVEIS PARA OS TEMPLATES ---
        current_date = datetime.now().strftime('%d/%m/%Y')

        output_dir = "/Users/cairorocha/Documents/fichas_tecnicas1/output"
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"FICHA_TECNICA_COMPLETA_{timestamp}.xlsx"
        filepath = os.path.join(output_dir, filename)

        template_vars = {
            'current_date': current_date,
            'local': 'Teresina, Piauí, Brasil',
            'filepath': filepath,
            'color': 'Blue'  # Cor personalizada para o Excel
        }

        # --- CRIAR TAREFAS A PARTIR DOS CONFIGS ---
        try:
            task_read_sources = create_task_from_config('task_read_sources', tasks_config['task_read_sources'], agents_dict, **template_vars)
            
            task_extract_fichas_tecnicas = create_task_from_config('task_extract_fichas_tecnicas', tasks_config['task_extract_fichas_tecnicas'], 
                                                                 agents_dict, {'task_read_sources': task_read_sources}, **template_vars)
            
            task_extract_base_insumos = create_task_from_config('task_extract_base_insumos', tasks_config['task_extract_base_insumos'], 
                                                               agents_dict, {'task_extract_fichas_tecnicas': task_extract_fichas_tecnicas}, **template_vars)
            
            task_consolidate_data = create_task_from_config('task_consolidate_data', tasks_config['task_consolidate_data'], 
                                                           agents_dict, {'task_extract_fichas_tecnicas': task_extract_fichas_tecnicas, 
                                                                        'task_extract_base_insumos': task_extract_base_insumos}, **template_vars)
            
            task_generate_excel = create_task_from_config('task_generate_excel', tasks_config['task_generate_excel'], 
                                                         agents_dict, {'task_consolidate_data': task_consolidate_data}, **template_vars)
            logger.info("Tasks created successfully")
        except Exception as e:
            logger.error(f"Error creating tasks: {e}")
            raise

        # --- MONTAGEM DA CREW ---
        final_crew = Crew(
            agents=[file_reader_agent, ficha_tecnica_agent, base_insumos_agent, data_consolidator_agent, excel_writer_agent],
            tasks=[task_read_sources, task_extract_fichas_tecnicas, task_extract_base_insumos, task_consolidate_data, task_generate_excel],
            process=Process.sequential,
            memory=False,
            verbose=False
        )
        logger.info("Crew assembled successfully")

        # --- EXECUÇÃO ---
        # Garantir que o diretório output existe
        output_dir = ensure_output_directory()
        
        # Usar fontes customizadas ou carregar da configuração
        if custom_sources:
            sources = custom_sources
            logger.info(f"Using custom sources: {len(sources)} provided")
        else:
            sources = sources_config.get('sources', [])
            if not sources:
                logger.error("No sources found in configuration")
                return
        
        # Validar fontes
        sources_validas = validate_sources(sources)
        
        if not sources_validas:
            logger.error("No valid sources found")
            return
        
        inputs = {'sources': sources_validas}
        logger.info(f"Starting processing of {len(sources_validas)} sources")
        logger.info(f"Valid sources: {sources_validas}")
        
        # Initialize Langfuse
        langfuse = init_langfuse()
        
        # Execute crew
        if langfuse:
            with langfuse.start_as_current_span(name="crewai-fichas-tecnicas-1"):
                result = final_crew.kickoff(inputs=inputs)
        else:
            result = final_crew.kickoff(inputs=inputs)

        logger.info("Process completed successfully")
        logger.info(f"Result: {result}")
        
        # Verificar arquivos gerados e limpar inválidos
        output_files = [f for f in os.listdir(output_dir) if f.endswith('.xlsx')]
        valid_files = []
        invalid_files = []
        
        for file in output_files:
            file_path = os.path.join(output_dir, file)
            if safe_file_check(file_path):
                file_size = os.path.getsize(file_path)
                logger.info(f"Valid Excel file created: {file_path} ({file_size} bytes)")
                valid_files.append(file_path)
            else:
                logger.warning(f"Invalid or corrupted file found: {file_path}")
                invalid_files.append(file_path)
        
        # Remove invalid files
        for invalid_file in invalid_files:
            try:
                os.remove(invalid_file)
                logger.info(f"Removed invalid Excel file: {invalid_file}")
            except Exception as e:
                logger.error(f"Failed to remove invalid file {invalid_file}: {e}")
        
        if valid_files:
            logger.info(f"Process completed successfully with {len(valid_files)} valid Excel file(s)")
            # Retornar o arquivo Excel mais recente
            latest_file = max(valid_files, key=os.path.getmtime)
            return {'success': True, 'excel_file': latest_file, 'result': result}
        else:
            logger.warning("No valid Excel files found in output directory")
            return {'success': False, 'excel_file': None, 'result': "No valid Excel files generated"}
        
    except Exception as e:
        logger.error(f"Error in fichas_tecnicas execution: {e}")
        raise


if __name__ == '__main__':
    try:
        fichas_tecnicas()
    except Exception as e:
        logger.error(f"Application failed: {e}")
        raise
